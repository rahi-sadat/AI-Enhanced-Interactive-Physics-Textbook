"""
scripts/generate_docs_bundle.py
Generates both:
1. A professionally styled HTML document (docs/AugmentedPhysics_v2_Comprehensive_Report.html)
2. A formal Word document (.docx) (docs/AugmentedPhysics_v2_Comprehensive_Report.docx)
Based on docs/COMPREHENSIVE_SYSTEM_REPORT_AND_BASE_PAPER_COMPARISON.md
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "docs"))
MD_PATH = os.path.join(DOCS_DIR, "COMPREHENSIVE_SYSTEM_REPORT_AND_BASE_PAPER_COMPARISON.md")
HTML_OUT = os.path.join(DOCS_DIR, "AugmentedPhysics_v2_Comprehensive_Report.html")
DOCX_OUT = os.path.join(DOCS_DIR, "AugmentedPhysics_v2_Comprehensive_Report.docx")

def read_markdown(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()

# Helper XML styling for docx tables
def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), val)
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            borders.append(border)
        tblPr[0].append(borders)

def build_docx(md_content, output_path):
    print("Generating DOCX:", output_path)
    doc = docx.Document()

    # Page setup: Normal Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Parse markdown line by line
    lines = md_content.split('\n')
    i = 0
    in_code = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines: return
        header_line = t_lines[0]
        headers = [c.strip() for c in header_line.split('|')[1:-1]]
        data_rows = []
        for r_line in t_lines[2:]:
            if '|' in r_line:
                cells = [c.strip() for c in r_line.split('|')[1:-1]]
                data_rows.append(cells)

        if not headers: return
        table = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table, "D0D5DD")

        # Header formatting
        hdr_cells = table.rows[0].cells
        for col_idx, h_text in enumerate(headers):
            cell = hdr_cells[col_idx]
            cell.text = re.sub(r'[*_`]', '', h_text)
            set_cell_background(cell, "0F2942") # Dark navy blue
            set_cell_margins(cell, 120, 120, 150, 150)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(10)

        # Body formatting
        for row_idx, row_data in enumerate(data_rows):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx in range(len(headers)):
                cell = row_cells[col_idx]
                val = row_data[col_idx] if col_idx < len(row_data) else ""
                clean_val = re.sub(r'[*_`]', '', val)
                cell.text = clean_val
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, 100, 100, 140, 140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(51, 65, 85)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)

    while i < len(lines):
        line = lines[i]

        # Handle code blocks
        if line.strip().startswith('```'):
            if in_code:
                in_code = False
                code_text = '\n'.join(code_lines)
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = Inches(0.3)
                code_p.paragraph_format.right_indent = Inches(0.3)
                code_p.paragraph_format.space_before = Pt(4)
                code_p.paragraph_format.space_after = Pt(6)
                run = code_p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if line.strip().startswith('|') and '|' in line:
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:].strip())
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 41, 66)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:].strip())
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:].strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif line.startswith('> '):
            # Blockquote
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            quote_text = re.sub(r'[*_`]', '', line[2:].strip())
            run = p.add_run(quote_text)
            run.font.italic = True
            run.font.color.rgb = RGBColor(75, 85, 99)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            bullet_text = re.sub(r'[*_`]', '', line[2:].strip())
            p.add_run(bullet_text)
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            num_text = re.sub(r'^\d+\.\s', '', line).strip()
            num_text = re.sub(r'[*_`]', '', num_text)
            p.add_run(num_text)
        elif line.strip() == '---':
            # Separator
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("____________________________________________________________________")
            run.font.color.rgb = RGBColor(203, 213, 225)
        elif line.strip() == '':
            pass
        else:
            p = doc.add_paragraph()
            clean_text = re.sub(r'[*_`]', '', line)
            p.add_run(clean_text)

        i += 1

    if in_table and table_lines:
        flush_table(table_lines)

    doc.save(output_path)
    print("DOCX successfully generated:", output_path)

def build_html(md_content, output_path):
    print("Generating HTML:", output_path)
    
    # Simple markdown parser for HTML
    html_body = []
    lines = md_content.split('\n')
    in_code = False
    code_lines = []
    in_table = False
    table_lines = []
    in_list = False

    def inline_fmt(txt):
        # bold
        txt = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', txt)
        # italic
        txt = re.sub(r'\*(.+?)\*', r'<em>\1</em>', txt)
        # code
        txt = re.sub(r'`(.+?)`', r'<code>\1</code>', txt)
        # links
        txt = re.sub(r'\[(.*?)\]\((.*?)\)', r'<a href="\2" target="_blank">\1</a>', txt)
        # Mathjax-like representation
        txt = re.sub(r'\$(.+?)\$', r'<span class="math">\1</span>', txt)
        return txt

    def flush_table_html(t_lines):
        if not t_lines: return ""
        out = ['<div class="table-container"><table><thead><tr>']
        headers = [c.strip() for c in t_lines[0].split('|')[1:-1]]
        for h in headers:
            out.append(f'<th>{inline_fmt(h)}</th>')
        out.append('</tr></thead><tbody>')
        for r_line in t_lines[2:]:
            if '|' in r_line:
                cells = [c.strip() for c in r_line.split('|')[1:-1]]
                out.append('<tr>')
                for c in cells:
                    out.append(f'<td>{inline_fmt(c)}</td>')
                out.append('</tr>')
        out.append('</tbody></table></div>')
        return '\n'.join(out)

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith('```'):
            if in_code:
                in_code = False
                code_str = '\n'.join(code_lines).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                html_body.append(f'<pre><code>{code_str}</code></pre>')
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith('|') and '|' in line:
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                html_body.append(flush_table_html(table_lines))
                table_lines = []

        if line.startswith('# '):
            html_body.append(f'<h1 class="doc-title">{inline_fmt(line[2:].strip())}</h1>')
        elif line.startswith('## '):
            html_body.append(f'<h2>{inline_fmt(line[3:].strip())}</h2>')
        elif line.startswith('### '):
            html_body.append(f'<h3>{inline_fmt(line[4:].strip())}</h3>')
        elif line.startswith('#### '):
            html_body.append(f'<h4>{inline_fmt(line[5:].strip())}</h4>')
        elif line.startswith('> '):
            html_body.append(f'<blockquote>{inline_fmt(line[2:].strip())}</blockquote>')
        elif line.startswith('- ') or line.startswith('* '):
            if not in_list:
                html_body.append('<ul>')
                in_list = True
            html_body.append(f'<li>{inline_fmt(line[2:].strip())}</li>')
        elif re.match(r'^\d+\.\s', line):
            if not in_list:
                html_body.append('<ol>')
                in_list = True
            content = re.sub(r'^\d+\.\s', '', line).strip()
            html_body.append(f'<li>{inline_fmt(content)}</li>')
        elif line.strip() == '---':
            if in_list:
                html_body.append('</ul>' if html_body[-1].startswith('<li>') else '</ol>')
                in_list = False
            html_body.append('<hr />')
        elif line.strip() == '':
            if in_list:
                html_body.append('</ul>' if '<ul>' in html_body[-2:] else '</ol>')
                in_list = False
        else:
            if in_list:
                html_body.append('</ul>' if '<ul>' in html_body[-2:] else '</ol>')
                in_list = False
            html_body.append(f'<p>{inline_fmt(line)}</p>')

        i += 1

    if in_table and table_lines:
        html_body.append(flush_table_html(table_lines))
    if in_list:
        html_body.append('</ul>')

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>AugmentedPhysics v2 — Comprehensive System Pipeline & Base Paper Comparison</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
  <style>
    :root {{
      --primary: #0284c7;
      --primary-dark: #0369a1;
      --accent: #38bdf8;
      --navy: #0f172a;
      --surface: #ffffff;
      --surface-subtle: #f8fafc;
      --border: #e2e8f0;
      --border-dark: #cbd5e1;
      --text: #1e293b;
      --text-muted: #64748b;
      --success: #10b981;
      --warning: #f59e0b;
      --code-bg: #f1f5f9;
    }}

    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
    }}

    body {{
      font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      color: var(--text);
      background-color: #f8fafc;
      line-height: 1.65;
      font-size: 15px;
      -webkit-font-smoothing: antialiased;
    }}

    .page-container {{
      max-width: 1040px;
      margin: 40px auto;
      background: var(--surface);
      border: 1px solid var(--border);
      border-radius: 12px;
      padding: 60px 80px;
      box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.05), 0 8px 10px -6px rgba(0, 0, 0, 0.03);
    }}

    .header-badge {{
      display: inline-block;
      padding: 4px 12px;
      background: rgba(56, 189, 248, 0.12);
      color: var(--primary-dark);
      font-weight: 600;
      font-size: 12px;
      border-radius: 9999px;
      text-transform: uppercase;
      letter-spacing: 0.05em;
      margin-bottom: 16px;
      border: 1px solid rgba(56, 189, 248, 0.3);
    }}

    h1.doc-title {{
      font-size: 28px;
      font-weight: 800;
      color: var(--navy);
      line-height: 1.25;
      margin-bottom: 12px;
      letter-spacing: -0.02em;
    }}

    h2 {{
      font-size: 20px;
      font-weight: 700;
      color: var(--navy);
      margin-top: 36px;
      margin-bottom: 14px;
      padding-bottom: 8px;
      border-bottom: 2px solid var(--border);
      letter-spacing: -0.01em;
    }}

    h3 {{
      font-size: 16px;
      font-weight: 600;
      color: var(--primary-dark);
      margin-top: 24px;
      margin-bottom: 10px;
    }}

    h4 {{
      font-size: 14px;
      font-weight: 600;
      color: #334155;
      margin-top: 18px;
      margin-bottom: 6px;
    }}

    p {{
      margin-bottom: 14px;
      color: #334155;
    }}

    ul, ol {{
      margin-left: 24px;
      margin-bottom: 16px;
      color: #334155;
    }}

    li {{
      margin-bottom: 6px;
    }}

    blockquote {{
      border-left: 4px solid var(--primary);
      padding: 12px 18px;
      background: var(--surface-subtle);
      border-radius: 0 8px 8px 0;
      font-style: italic;
      color: #475569;
      margin: 18px 0;
    }}

    hr {{
      border: none;
      border-top: 1px solid var(--border);
      margin: 32px 0;
    }}

    .table-container {{
      width: 100%;
      overflow-x: auto;
      margin: 20px 0 28px 0;
      border-radius: 8px;
      border: 1px solid var(--border);
      background: #ffffff;
    }}

    table {{
      width: 100%;
      border-collapse: collapse;
      font-size: 13.5px;
      text-align: left;
    }}

    th {{
      background: #0f172a;
      color: #ffffff;
      font-weight: 600;
      padding: 12px 16px;
      border: 1px solid #1e293b;
    }}

    td {{
      padding: 10px 16px;
      border: 1px solid var(--border);
      vertical-align: top;
      color: #334155;
    }}

    tr:nth-child(even) td {{
      background-color: #f8fafc;
    }}

    tr:hover td {{
      background-color: #f1f5f9;
    }}

    pre {{
      background: #0f172a;
      color: #f8fafc;
      padding: 16px 20px;
      border-radius: 8px;
      font-family: 'JetBrains Mono', Consolas, monospace;
      font-size: 12.5px;
      line-height: 1.5;
      overflow-x: auto;
      margin: 18px 0;
      box-shadow: inset 0 2px 4px rgba(0,0,0,0.2);
    }}

    code {{
      font-family: 'JetBrains Mono', Consolas, monospace;
      font-size: 12.5px;
      background: var(--code-bg);
      color: #0369a1;
      padding: 2px 6px;
      border-radius: 4px;
    }}

    pre code {{
      background: transparent;
      color: inherit;
      padding: 0;
    }}

    .math {{
      font-family: "Cambria Math", "Computer Modern", Georgia, serif;
      font-style: italic;
      color: #0f172a;
    }}

    .print-bar {{
      display: flex;
      justify-content: flex-end;
      gap: 12px;
      margin-bottom: 24px;
    }}

    .print-btn {{
      padding: 8px 16px;
      background: var(--primary);
      color: #ffffff;
      border: none;
      border-radius: 6px;
      font-weight: 600;
      font-size: 13px;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      gap: 6px;
      transition: background 0.15s ease;
    }}

    .print-btn:hover {{
      background: var(--primary-dark);
    }}

    @media print {{
      body {{
        background: #ffffff;
        font-size: 11pt;
      }}
      .page-container {{
        box-shadow: none;
        border: none;
        padding: 0;
        margin: 0;
        max-width: 100%;
      }}
      .print-bar {{
        display: none;
      }}
      h2 {{
        page-break-after: avoid;
      }}
      table, pre, blockquote {{
        page-break-inside: avoid;
      }}
    }}
  </style>
</head>
<body>
  <div class="page-container">
    <div class="print-bar">
      <button class="print-btn" onclick="window.print()">🖨️ Print / Save to PDF</button>
    </div>
    <div class="header-badge">Academic Research &amp; Engineering Report</div>
    {''.join(html_body)}
  </div>
</body>
</html>
"""

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_template)
    print("HTML successfully generated:", output_path)

if __name__ == "__main__":
    content = read_markdown(MD_PATH)
    build_docx(content, DOCX_OUT)
    build_html(content, HTML_OUT)
